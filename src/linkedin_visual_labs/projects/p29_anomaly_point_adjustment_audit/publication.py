"""Step 6 scholarly and local-release artifacts derived from frozen evidence only."""

# ruff: noqa: E501

from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .config import PROJECT_ID, repository_root
from .presentation import load_frozen_release, presentation_root, validate_presentation

TITLE = "Point Adjustment Under Random Scores: A Partial Replication and Finite-Sample Audit"
PUBLIC_CREATOR = "Keerthi Sagar Chegondi"
PUBLIC_VERSION = "0.1.0"
PUBLICATION_ROOT = presentation_root() / "publication"
METADATA_ROOT = presentation_root() / "publication_metadata"


@dataclass(frozen=True, slots=True)
class PublicationInputs:
    """Resolved frozen inputs used by publication generation."""

    release_data_path: Path
    claim_register_path: Path
    source_provenance_path: Path


def publication_inputs(input_root: Path | None = None) -> PublicationInputs:
    """Resolve canonical inputs or require a complete explicit clean-input bundle."""
    if input_root is None:
        return PublicationInputs(
            release_data_path=presentation_root() / "data" / "release_data.json",
            claim_register_path=presentation_root() / "manifests" / "claim_register.json",
            source_provenance_path=repository_root()
            / "reference"
            / "project9"
            / "source_provenance.json",
        )
    root = input_root.resolve()
    inputs = PublicationInputs(
        release_data_path=root / "data" / "release_data.json",
        claim_register_path=root / "manifests" / "claim_register.json",
        source_provenance_path=root / "reference" / "source_provenance.json",
    )
    for logical_name, path in (
        ("release_data", inputs.release_data_path),
        ("claim_register", inputs.claim_register_path),
        ("source_provenance", inputs.source_provenance_path),
    ):
        if not path.is_relative_to(root):
            raise ValueError(f"clean publication input escapes input_root: {logical_name}")
        if not path.is_file():
            raise FileNotFoundError(f"missing clean publication input {logical_name}: {path}")
    return inputs


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"expected JSON object: {path}")
    return value


def _f1(primary: dict[str, Any], key: str) -> float:
    value = primary[key]
    if not isinstance(value, dict) or not isinstance(value.get("mean"), int | float):
        raise ValueError(f"missing frozen result: {key}")
    return float(value["mean"])


def _evidence(inputs: PublicationInputs) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    release = load_frozen_release(inputs.release_data_path)
    claims = _load_json(inputs.claim_register_path)
    provenance = _load_json(inputs.source_provenance_path)
    if release.get("release_data_status") != "FROZEN" or not claims.get("approved_claims"):
        raise ValueError("publication requires frozen data and approved claims")
    return release, claims, provenance


def _manuscript_markdown(release: dict[str, Any]) -> str:
    primary = release["primary_condition"]
    if not isinstance(primary, dict):
        raise ValueError("missing primary condition")
    raw, adjusted = _f1(primary, "raw_f1"), _f1(primary, "adjusted_f1")
    return f"""# {TITLE}

**Author:** {PUBLIC_CREATOR}

## Abstract

This note is an independent partial replication and controlled synthetic audit of point adjustment, following the evaluation concern raised by Kim et al. It holds label-independent random alerts fixed and compares raw pointwise and point-adjusted scoring across predeclared event lengths and alert probabilities. In the frozen primary condition, mean F1 is {raw:.3f} under raw scoring and {adjusted:.3f} under point adjustment across 2,000 paired repetitions. This result is not a trained-detector benchmark, a full-paper reproduction, or a replacement metric proposal.

## Introduction

Evaluation design can change reported detector performance without changing detector predictions. Point adjustment credits an anomalous event when at least one prediction touches it. Kim et al. documented concerns about this protocol; Project 9 asks a narrower question: under an explicit synthetic Bernoulli alert model, how do event length and point adjustment alter reported paired scores?

## Background and Related Work

Kim et al. (2022) motivate the replication target. TADPAK provides reference implementation context only; no code was copied. Outliar provides overlapping public implementation context whose claims were not independently reproduced here. PATE provides proximity-aware evaluation context; Project 9 proposes no replacement metric.

## Point Adjustment Protocol

Raw pointwise evaluation credits each timestamp separately. Point adjustment retrospectively grants event-level credit after any alert touches a ground-truth event. The predictions are identical in both paired evaluations; only the scoring protocol changes.

## Replication Target

This is a partial replication of the mechanism and random-score behavior described by Kim et al., not a full reproduction of its datasets, detectors, benchmarks, or conclusions.

## Controlled Experimental Design

The frozen design uses N=20,000 timestamps, P=1,000 anomalous timestamps, prevalence 0.05, deterministic equal-event placement, event lengths 1, 5, 10, 20, 50, 100, 200, 500, and 1000, and q values 0.001, 0.005, 0.01, 0.05, and 0.10. The primary condition is L=200 and q=0.01. Each condition uses 2,000 paired repetitions, stable derived RNG streams, a sufficient-count simulator, and a timeline semantic oracle.

## Analytical Model

For an event of cardinality L and independent alert probability q, the event-hit probability is h(L,q) = 1-(1-q)^L. The release distinguishes expected counts and adjusted-recall expectation from finite-run mean F1. The ratio-of-expected-counts reference is not labeled expected finite-run F1.

## Finite-Sample Distinction

Monte Carlo mean F1 aggregates realized paired runs. Ratio references summarize expected counts. These quantities answer related but different questions and are retained separately in frozen release data.

## Results

Figure 1 shows that the same random alerts produce frozen primary-condition mean F1 values of {raw:.3f} raw and {adjusted:.3f} point-adjusted (approved claim P9-C001; `primary_condition.raw_f1.mean` and `primary_condition.adjusted_f1.mean`). Figure 2 shows q=0.01 event-length sensitivity. Figure 3 compares Monte Carlo adjusted recall with the frozen analytical event-hit expectation (approved claim P9-C002; `event_length_sweep_q_0_01`).

![Figure 1: Same alerts, different F1](../figures/01_same_alerts_different_f1.png)

![Figure 2: Event-length sensitivity](../figures/02_event_length_sensitivity.png)

![Figure 3: Probability validation](../figures/03_probability_validation.png)

## Validation

The release checks bounded metrics, paired F1 monotonicity, analytical raw counts, analytical adjusted recall, deterministic timelines, and fixed seeds. Presentation validation confirms the three figure pairs, six-tab dashboard, 45-second H.264 video, decoded keyframes, and social-draft package.

## Source Discrepancy

The source-as-printed equation ambiguity remains UNRESOLVED. Inclusive segment notation and endpoint-difference exponents create a cardinality ambiguity, so this project uses its own explicit Bernoulli generative model. It is not presented as a correction to the original paper.

## Discussion

The interpretation is limited: longer events give a random alert more opportunities to touch an event. This demonstrates metric behavior under the declared model, not detector quality or a real-world operational outcome.

## Limitations

The labels are synthetic and scores are label-independent. Version 1 evaluates no trained detector, is only a partial replication, does not establish that random scores beat modern systems, does not invalidate all event-aware metrics, and uses retrospective ground-truth event boundaries. Monte Carlo uncertainty applies only to this synthetic model. Editorial and venue scope remain to be reviewed.

## Reproducibility

The accompanying appendix specifies environment, configuration, seeds, condition grid, validation gates, hashes, and local commands. Frozen release data, rather than a different stochastic run, is the publication input.

## Contributions and AI Disclosure

The contribution is a transparent finite-sample audit with paired identical alerts and explicit distinction between expected-count ratios and finite-run F1. AI assistance supported coding, tests, literature-screen organization, drafting, and presentation production; the human author remains responsible for scientific interpretation and any future submission.

## References

Kim, S., Choi, K., Choi, H.-S., Lee, B., and Yoon, S. (2022). Towards a Rigorous Evaluation of Time-Series Anomaly Detection. AAAI.

Ghorbani, R., Reinders, M. J. T., and Tax, D. M. J. (2024). PATE: Proximity-Aware Time series anomaly Evaluation. KDD Research Track / arXiv.
"""


def _appendix(release: dict[str, Any], release_data_path: Path) -> str:
    return f"""# Project 9 Reproducibility Appendix

## Frozen evidence

Release input: `outputs/{PROJECT_ID}/data/release_data.json`{"  "}
SHA-256: `{_hash(release_data_path)}`

## Environment and configuration

Python: `{sys.version.split()[0]}`. Configuration: `configs/{PROJECT_ID}.yaml`. The model is CPU-only and uses NumPy stable derived seed streams.

## Design

N=20,000; P=1,000; prevalence=0.05; event lengths 1, 5, 10, 20, 50, 100, 200, 500, 1000; q values 0.001, 0.005, 0.01, 0.05, 0.10; 2,000 paired repetitions; primary L=200, q=0.01. Timelines are deterministic equal-event placements. Count simulation and a timeline semantic oracle provide complementary validation.

## Commands

`python -m linkedin_visual_labs anomaly-evaluation validate-inputs`{"  "}
`python -m linkedin_visual_labs anomaly-evaluation render-presentation`{"  "}
`python -m linkedin_visual_labs anomaly-evaluation build-publication`{"  "}
`python -m linkedin_visual_labs anomaly-evaluation validate-publication`

## Outputs and limitation

The frozen release drives figures, dashboard, video, manuscript, and release metadata. The source equation ambiguity remains UNRESOLVED; the project uses its independently stated Bernoulli model and does not claim a correction to Kim et al.
"""


def _bib(provenance: dict[str, Any]) -> str:
    sources = provenance.get("sources")
    if not isinstance(sources, list):
        raise ValueError("source provenance lacks sources")
    kim = next(source for source in sources if source.get("id") == "S1")
    pate = next(source for source in sources if source.get("id") == "S5")
    return f"""@inproceedings{{kim2022rigorous,
  title = {{{kim["title"]}}},
  author = {{Kim, Siwon and Choi, Kukjin and Choi, Hyun-Soo and Lee, Byunghan and Yoon, Sungroh}},
  year = {{2022}},
  doi = {{{kim["doi"]}}},
  url = {{{kim["url"]}}}
}}

@article{{ghorbani2024pate,
  title = {{{pate["title"]}}},
  author = {{Ghorbani, Ramin and Reinders, Marcel J. T. and Tax, David M. J.}},
  year = {{2024}},
  url = {{{pate["url"]}}}
}}
"""


def _make_docx(markdown: str, output: Path, evidence_root: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    for style_name in ("Normal", "Title", "Heading 1"):
        document.styles[style_name].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("!["):
            name = line.split("(")[1].rstrip(")")
            figure = evidence_root / name.replace("../", "")
            if figure.is_file():
                document.add_picture(str(figure), width=Inches(5.9))
        elif line:
            document.add_paragraph(line.replace("**", "").replace("`", ""))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def _make_pdf(markdown: str, output: Path, evidence_root: Path) -> None:
    from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore[import-untyped]
    from reportlab.platypus import (  # type: ignore[import-untyped]
        Image,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    styles = getSampleStyleSheet()
    story: list[Any] = []
    for line in markdown.splitlines():
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Title"]))
        elif line.startswith("## "):
            story.append(Spacer(1, 8))
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith("!["):
            figure = evidence_root / line.split("(")[1].rstrip(")").replace("../", "")
            if figure.is_file():
                story.append(Image(str(figure), width=420, height=263))
        elif line:
            story.append(
                Paragraph(
                    line.replace("&", "&amp;").replace("<", "&lt;").replace("**", ""),
                    styles["BodyText"],
                )
            )
    output.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(output), pagesize=letter, rightMargin=54, leftMargin=54).build(story)


def build_publication(
    root: Path | None = None,
    evidence_root: Path | None = None,
    input_root: Path | None = None,
) -> dict[str, str]:
    """Generate manuscript, appendix, metadata, and traceability from frozen inputs."""
    inputs = publication_inputs(input_root)
    release, claims, provenance = _evidence(inputs)
    markdown, appendix, bibliography = (
        _manuscript_markdown(release),
        _appendix(release, inputs.release_data_path),
        _bib(provenance),
    )
    target_root = (root or presentation_root()).resolve()
    publication_root = target_root / "publication"
    metadata_root = target_root / "publication_metadata"
    evidence = (evidence_root or presentation_root()).resolve()
    if root is not None and not publication_root.is_relative_to(target_root):
        raise ValueError("publication output escapes supplied root")
    publication_root.mkdir(parents=True, exist_ok=True)
    metadata_root.mkdir(parents=True, exist_ok=True)
    markdown_path = publication_root / "Point_Adjustment_Audit_Manuscript.md"
    docx_path = publication_root / "Point_Adjustment_Audit_Manuscript.docx"
    pdf_path = publication_root / "Point_Adjustment_Audit_Manuscript.pdf"
    markdown_path.write_text(markdown, encoding="utf-8")
    (publication_root / "references.bib").write_text(bibliography, encoding="utf-8")
    (publication_root / "reproducibility_appendix.md").write_text(appendix, encoding="utf-8")
    (publication_root / "AI_ASSISTANCE_LOG.md").write_text(
        "# AI Assistance Log\n\nAI assisted coding, test generation, literature-screen organization, manuscript drafting and editing, and figures/dashboard/video production. The human author remains responsible for scientific interpretation, author metadata, and any future submission.\n",
        encoding="utf-8",
    )
    _make_docx(markdown, docx_path, evidence)
    _make_pdf(markdown, pdf_path, evidence)
    citation = (
        f'cff-version: 1.2.0\ntitle: "{TITLE}"\ntype: software\nauthors:'
        f'\n  - name: "{PUBLIC_CREATOR}"\nversion: "{PUBLIC_VERSION}"\n'
    )
    (metadata_root / "CITATION.cff").write_text(citation, encoding="utf-8")
    (metadata_root / "zenodo.json").write_text(
        json.dumps(
            {
                "title": TITLE,
                "version": PUBLIC_VERSION,
                "creators": [{"name": PUBLIC_CREATOR}],
                "publication_date": None,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    status = {
        "project_id": PROJECT_ID,
        "manuscript_status": "READY_FOR_INTERNAL_REVIEW",
        "github_release": "NOT_PERFORMED",
        "zenodo": "NOT_PERFORMED",
        "arxiv": "NOT_PERFORMED",
        "journal": "NOT_PERFORMED",
        "portfolio": "NOT_PERFORMED",
        "linkedin": "NOT_PERFORMED",
        "citation_check": "NOT_CHECKED",
        "author_metadata": PUBLIC_CREATOR,
        "public_version": PUBLIC_VERSION,
    }
    (metadata_root / "publication_status.json").write_text(
        json.dumps(status, indent=2, sort_keys=True), encoding="utf-8"
    )
    trace = {
        "P9-C001": [
            "Results",
            "primary_condition.raw_f1.mean",
            "primary_condition.adjusted_f1.mean",
        ],
        "P9-C002": ["Results", "event_length_sweep_q_0_01"],
    }
    (metadata_root / "manuscript_claim_traceability.json").write_text(
        json.dumps(trace, indent=2, sort_keys=True), encoding="utf-8"
    )
    return {
        "markdown": str(markdown_path),
        "docx": str(docx_path),
        "pdf": str(pdf_path),
        "claims": str(len(claims["approved_claims"])),
    }


def validate_publication() -> dict[str, Any]:
    """Check the generated scholarly package without recalculating science."""
    required = [
        PUBLICATION_ROOT / name
        for name in (
            "Point_Adjustment_Audit_Manuscript.md",
            "Point_Adjustment_Audit_Manuscript.docx",
            "Point_Adjustment_Audit_Manuscript.pdf",
            "references.bib",
            "reproducibility_appendix.md",
            "AI_ASSISTANCE_LOG.md",
        )
    ]
    required += [
        METADATA_ROOT / name
        for name in (
            "CITATION.cff",
            "zenodo.json",
            "publication_status.json",
            "manuscript_claim_traceability.json",
        )
    ]
    missing = [str(path) for path in required if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise ValueError(f"missing publication artifacts: {missing}")
    from pypdf import PdfReader

    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(str(PUBLICATION_ROOT / "Point_Adjustment_Audit_Manuscript.pdf")).pages
    )
    required_headings = ("Abstract", "Results", "Source Discrepancy", "Limitations", "References")
    if not all(heading in text for heading in required_headings):
        raise ValueError("PDF text extraction lacks required manuscript headings")
    citation = (METADATA_ROOT / "CITATION.cff").read_text(encoding="utf-8")
    zenodo = _load_json(METADATA_ROOT / "zenodo.json")
    status = _load_json(METADATA_ROOT / "publication_status.json")
    metadata_text = "\n".join((citation, json.dumps(zenodo), json.dumps(status), text))
    if "AUTHOR_INPUT_REQUIRED" in metadata_text or "0.1.0-local" in metadata_text:
        raise ValueError("publication metadata contains a public placeholder")
    if (
        f'name: "{PUBLIC_CREATOR}"' not in citation
        or f'version: "{PUBLIC_VERSION}"' not in citation
    ):
        raise ValueError("CITATION.cff does not match the final public identity")
    if (
        zenodo.get("creators") != [{"name": PUBLIC_CREATOR}]
        or zenodo.get("version") != PUBLIC_VERSION
    ):
        raise ValueError("zenodo metadata does not match the final public identity")
    if "orcid" in zenodo or "doi" in zenodo:
        raise ValueError("zenodo metadata contains an unassigned persistent identifier")
    required_statuses = ("github_release", "zenodo", "arxiv", "journal", "portfolio", "linkedin")
    if any(status.get(name) != "NOT_PERFORMED" for name in required_statuses):
        raise ValueError("publication metadata overstates external publication status")
    return {
        "status": "PASS",
        "pdf_pages": len(
            PdfReader(str(PUBLICATION_ROOT / "Point_Adjustment_Audit_Manuscript.pdf")).pages
        ),
        "artifacts": len(required),
    }


def package_release() -> dict[str, str]:
    """Package the approved local-only deliverables and external checksum."""
    validate_publication()
    release_dir = presentation_root() / "release"
    release_dir.mkdir(parents=True, exist_ok=True)
    archive = release_dir / "project9_anomaly_evaluation_release_package.zip"
    include = [
        presentation_root() / "data",
        presentation_root() / "figures",
        presentation_root() / "dashboard",
        presentation_root() / "video",
        PUBLICATION_ROOT,
        METADATA_ROOT,
        presentation_root() / "manifests",
        repository_root() / "configs" / f"{PROJECT_ID}.yaml",
        repository_root() / "reference" / "project9",
    ]
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for item in include:
            if item.is_file():
                bundle.write(item, item.relative_to(repository_root()))
            else:
                for file in item.rglob("*"):
                    if file.is_file() and "test_render" not in file.parts:
                        bundle.write(file, file.relative_to(repository_root()))
    checksum = release_dir / f"{archive.name}.sha256"
    checksum.write_text(f"{_hash(archive)}  {archive.name}\n", encoding="ascii")
    manifest: dict[str, Any] = {
        "project_id": PROJECT_ID,
        "release_zip": str(archive),
        "release_zip_sha256": _hash(archive),
        "release_data_sha256": _hash(presentation_root() / "data" / "release_data.json"),
        "claim_register_sha256": _hash(presentation_root() / "manifests" / "claim_register.json"),
        "validation": {
            "presentation": validate_presentation(),
            "publication": validate_publication(),
        },
        "publication_status": "READY_FOR_INTERNAL_REVIEW",
    }
    manifest_path = release_dir / "release_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    return {
        "zip": str(archive),
        "sha256": manifest["release_zip_sha256"],
        "manifest": str(manifest_path),
    }
